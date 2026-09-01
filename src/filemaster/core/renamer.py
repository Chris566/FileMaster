"""重命名引擎.

W2 详细实现:
- 规划 (plan): 只生成结果, 不动文件
- 应用 (apply): 真实文件 IO, 支持 3 种冲突策略
- 撤销联动: 每个成功操作入 UndoStack
- 冲突策略: skip / overwrite / rename_new

W5 扩展:
- 命名空间占位符 {pdf_*} / {word_*} / {excel_*} / {image_*}
  (e.g. {pdf_title} / {word_paragraphs} / {excel_sheet_name} / {image_taken_at})
- apply_with_progress: 每文件后回调 on_progress

W7 扩展:
- apply_with_progress 接受 is_cancelled, 协作式取消 (文件之间检查)

W9 扩展:
- _apply_one 走 safe_rename, 硬中断 (Step A 后检查 cancel)
- ROLLBACK 状态不入 UndoStack (没真完成)
- cleanup_orphan_tmps 在 apply 入口清理 .tmp 残留
"""

from __future__ import annotations

import contextlib
import os
import re
import shutil
from collections.abc import Callable, Iterable
from dataclasses import dataclass
from datetime import datetime
from enum import Enum
from pathlib import Path

from filemaster.core.safe_rename import (
    SafeRenameResult,
    cleanup_orphan_tmps,
    safe_rename,
)
from filemaster.core.template import Template
from filemaster.core.undo import UndoEntry, UndoStack
from filemaster.utils.hash import file_hash


class ConflictStrategy(str, Enum):
    """冲突策略."""

    SKIP = "skip"           # 目标已存在则跳过该文件
    OVERWRITE = "overwrite"  # 目标已存在则覆盖（先备份到 undo）
    RENAME_NEW = "rename_new"  # 目标已存在则改名 (1) (2)...

    def __str__(self) -> str:
        return self.value


# 文件大小格式化
def format_size(size: int) -> str:
    """把字节数格式化为人类可读字符串.

    Args:
        size: 字节数
    Returns:
        如 "1.5 MB" / "512 B" / "2.3 GB"
    """
    if size < 0:
        return "0 B"
    units = ("B", "KB", "MB", "GB", "TB")
    value = float(size)
    for unit in units:
        if value < 1024 or unit == units[-1]:
            if unit == "B":
                return f"{int(value)} {unit}"
            return f"{value:.1f} {unit}"
        value /= 1024
    return f"{value:.1f} {units[-1]}"  # 不可达


def format_date(epoch: float, fmt: str = "%Y-%m-%d") -> str:
    """把 epoch 格式化为日期字符串.

    Args:
        epoch: Unix timestamp
        fmt: strftime 格式（默认 YYYY-MM-DD）
    Returns:
        格式化后的日期字符串
    """
    try:
        return datetime.fromtimestamp(epoch).strftime(fmt)
    except (OSError, ValueError, OverflowError):
        return ""


def get_excel_sheet_name(file: Path) -> str:
    """读取 Excel 文件的第一个 sheet 名（如果可能）.

    Args:
        file: Excel 文件路径
    Returns:
        第一个 sheet 名；非 Excel / 失败返回 ""
    """
    suffix = file.suffix.lower()
    if suffix not in (".xlsx", ".xlsm"):
        return ""
    wb = None
    try:
        from openpyxl import load_workbook  # type: ignore[import-untyped]

        # 注意：openpyxl 的 Workbook 不是 context manager，
        # 老的 with 用法会 AttributeError，需显式 close
        wb = load_workbook(file, read_only=True, data_only=True)
        # 优先 active.title，没有就取 worksheets[0]
        if wb.active and getattr(wb.active, "title", None):
            return wb.active.title
        if wb.worksheets:
            return wb.worksheets[0].title
        return ""
    except Exception:
        return ""
    finally:
        if wb is not None:
            with contextlib.suppress(Exception):
                wb.close()


@dataclass(frozen=True)
class RenameResult:
    """单个文件的重命名结果."""

    source: Path
    target: Path | None
    status: str  # OK | SKIPPED | CONFLICT | OVERWRITTEN | RENAMED | DRY_RUN | ERROR | ROLLBACK
    message: str = ""


# W5: 命名空间前缀
_PDF_PREFIX = "pdf_"
_WORD_PREFIX = "word_"
_EXCEL_PREFIX = "excel_"
_IMAGE_PREFIX = "image_"


def _namespaced_keys(prefix: str, names: Iterable[str]) -> set[str]:
    """构造一组带前缀的占位符名集合."""
    return {f"{prefix}{n}" for n in names}


_PDF_KEYS = _namespaced_keys(_PDF_PREFIX, [
    "title", "author", "subject", "pages", "created", "modified",
])
_WORD_KEYS = _namespaced_keys(_WORD_PREFIX, [
    "title", "author", "subject", "paragraphs", "created", "modified",
])
_EXCEL_KEYS = _namespaced_keys(_EXCEL_PREFIX, [
    "title", "author", "subject", "sheets", "sheet_name", "created", "modified",
])
_IMAGE_KEYS = _namespaced_keys(_IMAGE_PREFIX, [
    "width", "height", "taken_at", "camera_make", "camera_model",
    "format", "aspect_ratio",
])
# 扩展字段 (W3 已有的通用占位符, 走 metadata reader)
_METADATA_KEYS = {
    "Title", "Author", "Subject", "PageCount", "ImageWidth", "ImageHeight",
}


class Renamer:
    """重命名引擎.

    - plan(): 只生成 RenameResult，不动文件（dry-run）
    - apply(): 真实文件 IO，写 UndoStack（如提供）
    - apply_with_progress(): apply + 逐文件进度回调 + W7 取消 + W9 硬中断
    """

    def __init__(self, template: Template, prefix: str = "", start_index: int = 1) -> None:
        self._template = template
        self._prefix = prefix
        self._start_index = start_index
        self._index = start_index

    def reset_index(self) -> None:
        """重置序号到起始值."""
        self._index = self._start_index

    @property
    def template(self) -> Template:
        return self._template

    def _placeholders_used(self) -> set[str]:
        """返回模板中出现的占位符名集合（用于按需计算扩展 context）."""
        return {p.name for p in self._template.placeholders()}

    def _context_for(self, file: Path) -> dict[str, object]:
        """构造单个文件的占位符上下文.

        按模板实际使用的占位符按需计算昂贵的扩展字段（Hash / Sheet / Metadata）。
        """
        ctx: dict[str, object] = {
            "Prefix": self._prefix,
            "OriginalName": file.name,
            "BaseName": file.stem,
            "Extension": file.suffix.lstrip("."),
            "Index": self._index,
        }

        used = self._placeholders_used()
        need_stat = used & {"FileSize", "FileSizeBytes", "CreatedDate", "ModifiedDate"}
        if need_stat:
            try:
                stat = file.stat()
                if "FileSizeBytes" in used:
                    ctx["FileSizeBytes"] = stat.st_size
                if "FileSize" in used:
                    ctx["FileSize"] = format_size(stat.st_size)
                if "CreatedDate" in used:
                    ctx["CreatedDate"] = format_date(stat.st_ctime)
                if "ModifiedDate" in used:
                    ctx["ModifiedDate"] = format_date(stat.st_mtime)
            except OSError:
                ctx.setdefault("FileSizeBytes", 0)
                ctx.setdefault("FileSize", "0 B")
                ctx.setdefault("CreatedDate", "")
                ctx.setdefault("ModifiedDate", "")

        if used & {"HashShort", "Hash"}:
            try:
                digest = file_hash(file, "md5")
                if "HashShort" in used:
                    ctx["HashShort"] = digest[:8]
                if "Hash" in used:
                    ctx["Hash"] = digest
            except OSError:
                ctx.setdefault("HashShort", "")
                ctx.setdefault("Hash", "")

        if "Sheet" in used:
            ctx["Sheet"] = get_excel_sheet_name(file)

        # W3 + W5: 文档元数据
        metadata_keys = _METADATA_KEYS
        namespace_keys = used & (_PDF_KEYS | _WORD_KEYS | _EXCEL_KEYS | _IMAGE_KEYS)
        if used & metadata_keys or namespace_keys:
            from filemaster.core.metadata import MetadataReader
            try:
                meta = MetadataReader().read(file)
                # 通用占位符 (W3 兼容, 不动)
                if "Title" in used:
                    ctx["Title"] = meta.title
                if "Author" in used:
                    ctx["Author"] = meta.author
                if "Subject" in used:
                    ctx["Subject"] = meta.subject
                if "PageCount" in used:
                    ctx["PageCount"] = meta.page_count
                if "ImageWidth" in used or "ImageHeight" in used:
                    size = meta.extra.get("size", (0, 0)) if isinstance(meta.extra, dict) else (0, 0)
                    if "ImageWidth" in used:
                        ctx["ImageWidth"] = size[0]
                    if "ImageHeight" in used:
                        ctx["ImageHeight"] = size[1]
                # W5: 命名空间占位符
                if used & _PDF_KEYS:
                    if "pdf_title" in used:
                        ctx["pdf_title"] = meta.title
                    if "pdf_author" in used:
                        ctx["pdf_author"] = meta.author
                    if "pdf_subject" in used:
                        ctx["pdf_subject"] = meta.subject
                    if "pdf_pages" in used:
                        ctx["pdf_pages"] = meta.page_count
                    if "pdf_created" in used:
                        ctx["pdf_created"] = meta.created
                    if "pdf_modified" in used:
                        ctx["pdf_modified"] = meta.modified
                if used & _WORD_KEYS:
                    if "word_title" in used:
                        ctx["word_title"] = meta.title
                    if "word_author" in used:
                        ctx["word_author"] = meta.author
                    if "word_subject" in used:
                        ctx["word_subject"] = meta.subject
                    if "word_paragraphs" in used:
                        ctx["word_paragraphs"] = meta.paragraphs
                    if "word_created" in used:
                        ctx["word_created"] = meta.created
                    if "word_modified" in used:
                        ctx["word_modified"] = meta.modified
                if used & _EXCEL_KEYS:
                    if "excel_title" in used:
                        ctx["excel_title"] = meta.title
                    if "excel_author" in used:
                        ctx["excel_author"] = meta.author
                    if "excel_subject" in used:
                        ctx["excel_subject"] = meta.subject
                    if "excel_sheets" in used:
                        ctx["excel_sheets"] = meta.sheets_count
                    if "excel_sheet_name" in used:
                        ctx["excel_sheet_name"] = get_excel_sheet_name(file)
                    if "excel_created" in used:
                        ctx["excel_created"] = meta.created
                    if "excel_modified" in used:
                        ctx["excel_modified"] = meta.modified
                if used & _IMAGE_KEYS:
                    if "image_width" in used:
                        ctx["image_width"] = meta.width
                    if "image_height" in used:
                        ctx["image_height"] = meta.height
                    if "image_taken_at" in used:
                        ctx["image_taken_at"] = meta.taken_at
                    if "image_camera_make" in used:
                        ctx["image_camera_make"] = meta.camera_make
                    if "image_camera_model" in used:
                        ctx["image_camera_model"] = meta.camera_model
                    if "image_format" in used:
                        ctx["image_format"] = meta.image_format
                    if "image_aspect_ratio" in used:
                        ctx["image_aspect_ratio"] = meta.aspect_ratio
            except Exception:
                for k in used & metadata_keys:
                    if k in ("PageCount", "ImageWidth", "ImageHeight"):
                        ctx.setdefault(k, 0)
                    else:
                        ctx.setdefault(k, "")
                for k in used & namespace_keys:
                    if k.endswith(("_sheets", "_paragraphs", "_width", "_height", "_pages")):
                        ctx.setdefault(k, 0)
                    else:
                        ctx.setdefault(k, "")

        # W4 v1: 分类（lazy import — classifier 启动 0 成本但保持一致模式）
        category_keys = {"Category", "Category_zh"}
        if used & category_keys:
            from filemaster.core.classifier import classify_file as _cf
            try:
                c = _cf(file)
                if "Category" in used:
                    ctx["Category"] = c.category.value
                if "Category_zh" in used:
                    ctx["Category_zh"] = c.category.label_zh
            except Exception:
                ctx.setdefault("Category", "UNKNOWN")
                ctx.setdefault("Category_zh", "未知")

        return ctx

    def _render_target(self, file: Path) -> Path | None:
        """根据模板 + context 渲染出新文件名.

        Returns:
            新文件路径；如果渲染结果与原名相同则返回 None
        """
        ctx = self._context_for(file)
        new_name = self._template.render(ctx)
        if not new_name or new_name == file.name:
            return None
        # 用 sanitize 兜底
        new_name = self.sanitize(new_name)
        return file.with_name(new_name)

    def plan(self, files: Iterable[Path]) -> list[RenameResult]:
        """规划：只生成结果，不实际改文件.

        Args:
            files: 源文件列表
        Returns:
            RenameResult 列表（status="DRY_RUN" 或 "SKIPPED"）
        """
        results: list[RenameResult] = []
        for file in files:
            target = self._render_target(file)
            if target is None:
                results.append(RenameResult(file, None, "SKIPPED", "模板未变"))
            else:
                results.append(RenameResult(file, target, "DRY_RUN"))
            self._index += 1
        return results

    def apply(
        self,
        files: Iterable[Path],
        conflict_strategy: ConflictStrategy = ConflictStrategy.SKIP,
        undo_stack: UndoStack | None = None,
        is_cancelled: Callable[[], bool] | None = None,
    ) -> list[RenameResult]:
        """真实执行重命名.

        Args:
            files: 源文件列表
            conflict_strategy: 冲突策略（skip / overwrite / rename_new）
            undo_stack: 可选，提供则把每次成功操作写入栈
            is_cancelled: W7 协作式取消. W9 硬中断走 _apply_one, 此处也透传.
        Returns:
            RenameResult 列表
        """
        # W9: 入口清理 .tmp 残留 (上轮取消 / 崩溃遗留)
        self._cleanup_tmps(files)

        results: list[RenameResult] = []
        entries: list[UndoEntry] = []

        for file in files:
            # W7 协作式取消 — 文件之间检查
            if is_cancelled is not None and is_cancelled():
                break
            result, entry = self._apply_one(file, conflict_strategy, undo_stack, is_cancelled)
            results.append(result)
            if entry is not None:
                entries.append(entry)

        if undo_stack is not None and entries:
            undo_stack.push(entries)

        return results

    def apply_with_progress(
        self,
        files: Iterable[Path],
        conflict_strategy: ConflictStrategy = ConflictStrategy.SKIP,
        undo_stack: UndoStack | None = None,
        on_progress: Callable[[int, int, Path, RenameResult], None] | None = None,
        is_cancelled: Callable[[], bool] | None = None,
    ) -> list[RenameResult]:
        """W5: apply + 逐文件进度回调.

        W7 协作式取消 (文件之间检查).
        W9 硬中断 (_apply_one 内部 Step A 后检查).

        与 apply 行为一致, 但每处理完一个文件调 on_progress(current_index, total, file, result).
        on_progress 不抛异常, 内部吞掉.

        Args:
            files: 源文件列表
            conflict_strategy: 冲突策略
            undo_stack: 可选, 提供则写撤销栈
            on_progress: 回调 (i, total, file, result) -> None
            is_cancelled: 可选取消回调, 返回 True 时停止处理剩余文件.
                          W7: 文件之间协作式检查.
                          W9: _apply_one 内部 Step A 后硬中断检查.
                          已收集的 results 仍返回, 已入栈的 entries 保留;
                          ROLLBACK 状态的 file 不入栈 (没真完成).
        Returns:
            RenameResult 列表
        """
        # W9: 入口清理 .tmp 残留
        self._cleanup_tmps(files)

        files_list = list(files)
        total = len(files_list)
        results: list[RenameResult] = []
        entries: list[UndoEntry] = []

        for i, file in enumerate(files_list, 1):
            # W7 协作式取消 — 在文件之间检查
            if is_cancelled is not None and is_cancelled():
                break
            result, entry = self._apply_one(file, conflict_strategy, undo_stack, is_cancelled)
            results.append(result)
            if entry is not None:
                entries.append(entry)
            if on_progress is not None:
                with contextlib.suppress(Exception):
                    on_progress(i, total, file, result)

        if undo_stack is not None and entries:
            undo_stack.push(entries)

        return results

    def _apply_one(
        self,
        file: Path,
        conflict_strategy: ConflictStrategy,
        undo_stack: UndoStack | None,
        is_cancelled: Callable[[], bool] | None = None,
    ) -> tuple[RenameResult, UndoEntry | None]:
        """单文件 apply, 返回 (result, undo_entry).

        W9: 走 safe_rename. 取消时返回 ROLLBACK 状态, 不写 undo entry.
        """
        target = self._render_target(file)
        if target is None:
            self._index += 1
            return RenameResult(file, None, "SKIPPED", "模板未变"), None

        if target.exists():
            if conflict_strategy is ConflictStrategy.SKIP:
                self._index += 1
                return (
                    RenameResult(file, target, "CONFLICT", f"目标已存在：{target.name}（已跳过）"),
                    None,
                )
            if conflict_strategy is ConflictStrategy.RENAME_NEW:
                target = self._find_free_name(target)
                if target is None:
                    self._index += 1
                    return (
                        RenameResult(file, None, "ERROR", f"无法找到空闲名：{file.name}"),
                        None,
                    )
                final_status = "RENAMED"
                final_msg = f"已避开冲突：{target.name}"
            else:  # OVERWRITE
                final_status = "OVERWRITTEN"
                final_msg = f"已覆盖：{target.name}"
        else:
            final_status = "OK"
            final_msg = ""

        # W9: OVERWRITE 策略需要先备份被覆盖的目标 — 但 safe_rename 不直接做
        # 走两步: (1) 备份 target (2) safe_rename(file -> target)
        # 如 (1) 后取消, target 还在原位 (没动), 备份已写需清理
        backup_path: Path | None = None
        if (
            conflict_strategy is ConflictStrategy.OVERWRITE
            and target.exists()
            and undo_stack is not None
        ):
            backup_path = UndoStack.backup(target, self._backup_dir(undo_stack))

        # W9: safe_rename 内部 Step A 后检查 cancel, 取消时 ROLLBACK (源文件回原位)
        safe_result: SafeRenameResult = safe_rename(file, target, is_cancelled)

        if safe_result.status == "ROLLBACK":
            # 取消: 源文件已在原位, 不写 undo entry
            # 注意: backup_path 已写 (如果走了 OVERWRITE), 但因为源文件没动, target 也没被覆盖
            # → backup 是个孤儿, 需要清理
            if backup_path is not None:
                with contextlib.suppress(OSError):
                    backup_path.unlink()
            self._index += 1
            return RenameResult(file, target, "ROLLBACK", "已取消, 源文件保留"), None

        if safe_result.status == "ERROR":
            # safe_rename 失败: 源文件可能已动, 残留 .tmp
            # OVERWRITE 走过的备份也是孤儿, 清掉
            if backup_path is not None:
                with contextlib.suppress(OSError):
                    backup_path.unlink()
            self._index += 1
            return RenameResult(file, target, "ERROR", safe_result.message), None

        # OK: rename 成功
        self._index += 1
        entry: UndoEntry | None = None
        if undo_stack is not None:
            entry = UndoEntry(
                operation="RenameOnly",
                source=file,
                target=target,
                backup_path=backup_path,
            )
        return RenameResult(file, target, final_status, final_msg), entry

    @staticmethod
    def _cleanup_tmps(files: Iterable[Path]) -> int:
        """W9: 清理源文件所在目录的 .filemaster.tmp.* 残留.

        入口调用 (apply / apply_with_progress 开头), 应对:
        - 上轮取消时 safe_rename rollback 失败
        - 进程在 Step A / Step B 之间被杀
        - 用户手动 kill -9
        """
        cleaned = 0
        seen_dirs: set[Path] = set()
        for f in files:
            d = f.parent if f.is_file() or not f.exists() else f
            if d in seen_dirs:
                continue
            seen_dirs.add(d)
            cleaned += cleanup_orphan_tmps(d)
        return cleaned

    @staticmethod
    def _backup_dir(undo_stack: UndoStack) -> Path:
        """根据 UndoStack 推断备份目录."""
        if undo_stack._persist_dir is not None:
            return undo_stack._persist_dir / "backups"
        # 兜底：用户态临时
        import tempfile

        return Path(tempfile.gettempdir()) / "filemaster_backups"

    @staticmethod
    def _find_free_name(target: Path) -> Path | None:
        """找一个空闲的新名：name (1).ext / name (2).ext / ...

        上限 9999，避免无限循环。
        """
        stem, ext = target.stem, target.suffix
        for i in range(1, 10_000):
            candidate = target.with_name(f"{stem} ({i}){ext}")
            if not candidate.exists():
                return candidate
        return None

    def already_has_prefix(self, file: Path) -> bool:
        """判断文件是否已带前缀（不重复加）."""
        if not self._prefix:
            return False
        return file.stem.lower().startswith(self._prefix.lower())

    @staticmethod
    def sanitize(name: str) -> str:
        """去除 Windows 非法字符.

        Args:
            name: 原始文件名
        Returns:
            清理后的文件名
        """
        # Windows 非法字符：<>:"/\\|?*，以及控制字符；尾部空格/点也清掉
        cleaned = re.sub(r'[<>:"/\\|?*\x00-\x1f]', "_", name)
        return cleaned.rstrip(" .")
